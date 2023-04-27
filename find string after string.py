from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table, _Row
from docx.text.paragraph import Paragraph
import docx
import pandas as pd
import glob
import os

path = r'C:\Users\User\Desktop\Рабочая папка\от Геращенковой\2023\февраль 2023\РПД_2018_10.05.04_аид_О\*.docx'

g_glob = glob.glob(path)


df_tables_list = []

stop_kaf = ['ФВиС', 'ГиСД', 'ВМ', 'ОФ', 'ИнЯз']

def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    elif isinstance(parent, _Row):
        parent_elm = parent._tr
    else:
        raise ValueError("Что-то пошло не так!!")
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

for g_file in g_glob:
    g_file_short = g_file.replace(r'C:\Users\User\Desktop\Рабочая папка\от Геращенковой\2023\февраль 2023\РПД_2018_10.05.04_аид_О' + '\\', '')
    dir_name = os.path.dirname(g_file).replace(r'C:\Users\User\Desktop\Рабочая папка\от Геращенковой\2023\февраль 2023' + '\\', '')
    # print(dir_name)
    if g_file_short.startswith('РПД_'):
        doc = docx.Document(g_file)
        for block in iter_block_items(doc):
            # read Paragraph
            if isinstance(block, Paragraph):
                # print('ЭТО РПД')
                print(block.text)
            # read table
            elif isinstance(block, Table):
                print(block.table.cell(0,0).text)
                print(block.table.rows)
                print(g_file_short, doc)
                if block.table.cell(0,0).text == 'РАБОЧАЯ ПРОГРАММА учебной дисциплины':
                    print(block.table.cell(0,0).text)
                    
#                     df = [['' for i in range(len(block.table.columns))] for j in range(len(block.table.rows))]
#                     for r, row in enumerate(block.table.rows):
#                         for c, cell in enumerate(row.cells):
#                             # print(r, c, cell.text)
#                             df[r][c] = cell.text
#                     df_tables = pd.DataFrame(df)
#                     df_tables['Имя_файла'] = g_file_short
#                     # print(df)
#                     df_tables_list.append(df_tables)

# df_concat = pd.concat(df_tables_list)

# print(df_tables)
# print(df_concat)

# df_concat.to_excel(f'Сборка_{dir_name}.xlsx')
                      
        
    

# Таблица 1 – Планируемые результаты обучения по дисциплине
# Таблица 2 – Распределение трудоемкости дисциплины по видам учебной работы и семестрам
# Таблица 3 – Тематический план дисциплины
# Таблица 4 – Формирование компетенций по разделам дисциплины
# Таблица 5 – Тематика и содержание лекций
# Таблица 6 – Тематика лабораторных работ
# Таблица 7 – Тематика и содержание практических занятий
# Таблица 8 – Вопросы для самостоятельного изучения дисциплины
# Таблица 9 – Виды самостоятельной работы
# Таблица 10 – Формы и периодичность текущего контроля успеваемости
# Таблица 11 – Образовательные технологии, применяемые в ходе преподавания дисциплины
# Таблица 12 – Методические рекомендации обучающимся по освоению дисциплины
# Таблица 13 – Виды и средства оценивания результатов освоения дисциплины
# Таблица 14 – Критерии и шкала оценки РГР / доклада (реферата), его пре-зентации (выбрать необходимое) по дисциплине
# Таблица 15 – Шкала оценивания при промежуточной аттестации обучающихся
# Таблица 16 – Шкала оценивания, применяемая при выполнении и защите курсовой работы (курсового проекта) для технических дисциплин
# Таблица 17 – Шкала оценивания, применяемая при выполнении и защите курсовой работы (курсового проекта) для гуманитарных дисциплин
# Таблица 18 – Характеристика результатов обучения по дисциплине